"""El número de órganos del cuestionario se refleja en el articulado."""
import os
import sys
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from processors.estatutos import generar_estatutos  # noqa: E402

TMPL = os.path.join(os.path.dirname(__file__), "plantillas", "estatutos_template.docx")


def persona(n):
    return {"nombre": f"Persona Test {n}", "tipo_doc": "CC", "id_num": f"9999999{n}"}


def base(**extra):
    d = {
        "nombre_sas": "PRUEBA S.A.S.", "fecha": date(2026, 9, 25),
        "municipio": "Medellín", "departamento": "Antioquia",
        "accionistas": [{"tipo": "natural", "nombre": "Persona Test", "id_num": "99999999",
                         "expedicion": "Medellín", "domicilio": "Medellín",
                         "genero": "M", "porcentaje": 100}],
        "objeto_social": "cualquier actividad comercial o civil lícita",
        "capital_suscrito": 1_000_000, "capital_pagado": 1_000_000,
        "rl_principales": [{"nombre": "Persona Test", "cc": "99999999",
                            "tipo_doc": "CC", "expedicion": "Medellín"}],
        "rl_suplentes": [],
    }
    d.update(extra)
    return d


def generar(data):
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "e.docx")
        generar_estatutos(data, TMPL, out)
        doc = Document(out)
        return doc, "\n".join(p.text for p in doc.paragraphs)


# ── Sin junta ni suplentes: articulado original ──
doc, txt = generar(base())
assert "CAPÍTULO V - Representación legal" in txt
assert "Junta directiva" not in txt.split("CAPÍTULO X")[0] or "CAPÍTULO V - Junta" not in txt
assert "podrá tener uno o varios suplentes" in txt
assert "con su respectivo suplente," in txt

# ── Junta 3+2, dos RL suplentes, revisor sin suplente ──
doc, txt = generar(base(
    junta_directiva={"principales": [persona(1), persona(2), persona(3)],
                     "suplentes": [persona(4), persona(5)]},
    rl_suplentes=[{"nombre": "Persona Test 6", "cc": "1", "tipo_doc": "CC", "expedicion": "Cali"},
                  {"nombre": "Persona Test 7", "cc": "2", "tipo_doc": "CC", "expedicion": "Cali"}],
    revisor_fiscal={"tipo": "natural", "nombre": "Persona Test 8", "id_num": "3",
                    "tarjeta_profesional": "1-T"},
))
assert "La junta directiva; y" in txt and "La asamblea general de accionistas;\n" in txt
assert "asamblea general de accionistas y a la junta directiva," in txt
assert "CAPÍTULO V - Junta directiva" in txt
assert "CAPÍTULO VI - Representación legal" in txt and "CAPÍTULO XI - Disposiciones transitorias" in txt
assert "integrada por tres (3) miembros principales y dos (2) suplentes personales" in txt
assert "El representante legal tendrá dos (2) suplentes, designados" in txt
assert "un revisor fiscal, quien podrá tener un suplente," in txt
assert "tres (3) miembros principales y dos (2) miembros suplentes personales" in txt
# Coherencia del resto del articulado con la junta ya creada
assert "si se llegare a crear este órgano" not in txt
assert "Elegir y remover libremente a los miembros de la junta directiva" in txt
assert "salvo en la elección de la junta directiva o de otros cuerpos colegiados" in txt
assert "En ningún caso los accionistas podrán fraccionar su voto." not in txt

# Los artículos nuevos copian el formato del artículo de referencia
ps = doc.paragraphs
ref = next(p for p in ps if p.text.startswith("Nombramiento y período del representante legal."))
nuevo = next(p for p in ps if p.text.startswith("Composición y período de la junta directiva."))
assert nuevo._p.find(qn("w:pPr")).xml == ref._p.find(qn("w:pPr")).xml
assert [r.bold for r in nuevo.runs] == [r.bold for r in ref.runs]
assert nuevo.runs[1]._r.find(qn("w:rPr")).xml == ref.runs[1]._r.find(qn("w:rPr")).xml \
    if ref.runs[1]._r.find(qn("w:rPr")) is not None else True

# ── Junta 1+1, revisor con suplente ──
doc, txt = generar(base(
    junta_directiva={"principales": [persona(1)], "suplentes": [persona(2)]},
    revisor_fiscal={"tipo": "natural", "nombre": "Persona Test 8", "id_num": "3",
                    "tarjeta_profesional": "1-T",
                    "suplente": {"nombre": "Persona Test 9", "tipo_doc": "CC",
                                 "id_num": "4", "tarjeta_profesional": "2-T"}},
    rl_suplentes=[{"nombre": "Persona Test 6", "cc": "1", "tipo_doc": "CC", "expedicion": "Cali"}],
))
assert "integrada por un (1) miembro principal, con su respectivo suplente personal" in txt
assert "El representante legal tendrá un (1) suplente, designado por la asamblea general de accionistas, el cual tendrá" in txt
assert "con su respectivo suplente," in txt
assert "Revisor fiscal suplente: PERSONA TEST 9" in txt

print("OK test_organos")
