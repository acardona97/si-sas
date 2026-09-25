"""Módulo familia: del cuestionario a los estatutos, sin campos sin llenar."""
import os
import sys
import tempfile
from datetime import date

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "src"))
from docx import Document  # noqa: E402
from processors.familia import (  # noqa: E402
    construir_contexto, escenarios_por_defecto, DatosFamiliaError)
from processors.plantilla_familia import renderizar  # noqa: E402

TMPL = os.path.join(BASE, "plantillas", "estatutos_familia_template.docx")

DATA = {
    "nombre_sas": "FAMILIA PRUEBA S.A.S.", "municipio": "Envigado", "departamento": "Antioquia",
    "camara_ciudad": "Aburrá Sur",
    "capital_autorizado": 100_000_000, "capital_suscrito": 10_000_000, "valor_nominal": 1000,
    "accionistas": [
        {"tipo": "natural", "nombre": "Persona Test Uno", "tipo_doc": "CC", "id_num": "99999991",
         "domicilio": "Envigado", "porcentaje": 60, "clase": "A", "capital_pagado_num": 6_000_000},
        {"tipo": "natural", "nombre": "Persona Test Dos", "tipo_doc": "CC", "id_num": "99999992",
         "domicilio": "Medellín", "porcentaje": 30, "clase": "B", "capital_pagado_num": 1_000_000},
        {"tipo": "juridica", "nombre": "Holding Prueba S.A.S.", "id_num": "900111222-3",
         "rl_nombre": "Persona Test Tres", "rl_cc": "99999993", "porcentaje": 10, "clase": "B"},
    ],
    "rl_principales": [{"nombre": "Persona Test Uno", "cc": "99999991", "tipo_doc": "CC",
                        "expedicion": "Envigado"}],
    "rl_suplentes": [{"nombre": "Persona Test Cuatro", "cc": "99999994", "tipo_doc": "CC",
                      "expedicion": "Medellín"}],
    "opciones": {"junta_directiva": False, "consejo_familia": True, "protocolo_familia": True,
                 "arbitraje": True, "prohibicion_temporal": True, "autorizacion_transferencias": True,
                 "conversion_por_transferencia": True},
    "elecciones": {"grupo_familiar": "amplio", "preferencia": "simple", "prohibicion_anios": 15},
}


def generar(data):
    out = os.path.join(tempfile.mkdtemp(), "fam.docx")
    renderizar(TMPL, construir_contexto(data, hoy=date(2026, 9, 25)), out)
    doc = Document(out)
    return doc, "\n".join(p.text for p in doc.paragraphs)


doc, txt = generar(DATA)
# Objeto fijo de precautelación patrimonial
assert "la tenencia, precautelación e inversión de activos" in txt
# Negativas de opción de compra y exclusión en esta versión
assert "No se pacta opción estatutaria de compra" in txt
assert "No se pactan causales estatutarias de exclusión" in txt
# Decisiones: 70 % para gravámenes; excepciones condicionadas al grupo familiar
assert "setenta por ciento (70 %)" in txt
assert "el beneficiario final continúe perteneciendo al grupo familiar" in txt
# La prohibición de negociar nunca excede diez años aunque se pidan quince
assert "durante 10 años" in txt and "durante 15 años" not in txt
# Escenarios A, B y A+B
assert "Escenario A+B" in txt
# Suscripciones: 3 filas + encabezado + total; el PJ firma por su representante
filas = [[c.text for c in r.cells] for r in doc.tables[0].rows]
assert len(filas) == 5 and filas[-1][2] == "10.000", filas
assert "PERSONA TEST TRES" in txt and "HOLDING PRUEBA S.A.S." in txt
# El suplente, que no es accionista, firma aceptando el cargo
assert "acepta el cargo de representante legal suplente" in txt
assert "{{" not in txt and "[[" not in txt

# Una clase: sin conversión ni decisiones reservadas
una = dict(DATA, clases=[{"codigo": "A", "denominacion": "Ordinarias", "votos_por_accion": 1}],
           accionistas=[dict(a, clase="A") for a in DATA["accionistas"]])
_, txt1 = generar(una)
assert "Escenario A+B" not in txt1

# Validaciones del cuestionario
assert len(escenarios_por_defecto(["A", "B", "C"])) == 7
for malo, msg in [
    (dict(DATA, escenarios={"A": {"A": 100}}), "escenarios"),
    (dict(DATA, accionistas=[dict(DATA["accionistas"][0], porcentaje=50)]), "100 %"),
    (dict(DATA, elecciones={"grupo_familiar": "personalizado"}), "personalizada"),
]:
    try:
        construir_contexto(malo)
        raise AssertionError("debía fallar: " + msg)
    except DatosFamiliaError as e:
        assert msg in str(e), e

print("OK test_familia")
