# -*- coding: utf-8 -*-
"""Integración del módulo familia con el paquete y las validaciones del backend."""
import io
import os
import zipfile
from copy import deepcopy
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader

from test_paquete import _cliente, _payload, PAYLOAD_B, flask_app


# Mismo escenario de test_familia: dos personas naturales y una jurídica.
DATA = {
    "nombre_sas": "FAMILIA PRUEBA S.A.S.", "municipio": "Envigado", "departamento": "Antioquia",
    "camara_ciudad": "Aburrá Sur",
    "capital_autorizado": 100_000_000, "capital_suscrito": 10_000_000, "valor_nominal": 1000,
    "accionistas": [
        {"tipo": "natural", "nombre": "Persona Test Uno", "tipo_doc": "CC", "id_num": "99999991",
         "domicilio": "Envigado", "porcentaje": 60, "clase": "A", "capital_pagado_num": 6_000_000},
        {"tipo": "natural", "nombre": "Persona Test Dos", "tipo_doc": "CC", "id_num": "99999992",
         "domicilio": "Medellín", "porcentaje": 30, "clase": "B", "capital_pagado_num": 1_000_000},
        {"tipo": "juridica", "nombre": "Holding Prueba S.A.S.", "id_num": "900111222-3",
         "rl_nombre": "Persona Test Tres", "rl_cc": "99999993", "porcentaje": 10, "clase": "B"},
    ],
    "rl_principales": [{"nombre": "Persona Test Uno", "cc": "99999991", "tipo_doc": "CC",
                        "expedicion": "Envigado"}],
    "rl_suplentes": [{"nombre": "Persona Test Cuatro", "cc": "99999994", "tipo_doc": "CC",
                      "expedicion": "Medellín"}],
    "opciones": {"junta_directiva": False, "consejo_familia": True, "protocolo_familia": True,
                 "arbitraje": True, "prohibicion_temporal": True, "autorizacion_transferencias": True,
                 "conversion_por_transferencia": True},
    "elecciones": {"grupo_familiar": "amplio", "preferencia": "simple", "prohibicion_anios": 15},
}


def payload_familia():
    """Datos familiares con los campos comerciales que necesitan los anexos."""
    data = deepcopy(_payload("payload_b.json", PAYLOAD_B))
    # El cuestionario comercial incluye la clave aunque no haya disposiciones.
    data.pop("disposiciones", None)
    data.pop("modelo_propio", None)
    data.update(deepcopy(DATA))
    data.update(modulo="familia", es_empresa_familiar=False, nucleo_familiar=[])
    for acc in data["accionistas"]:
        if acc["tipo"] == "natural":
            acc["parentesco"] = "Padre" if acc["clase"] == "A" else "Hija"
        # El paquete recibe los aportes individuales como los envía el cuestionario.
        acc["capital_pagado"] = str(acc.get("capital_pagado_num", 1_000_000))
    return data


def test_paquete_familia(client):
    data = payload_familia()
    data.update(ciiu_code="5611", objeto_social="Texto que debe ignorarse",
                ciiu_code_sec="6201", ciiu_description_sec="Desarrollo de software")
    with patch.object(flask_app, "generar_objeto_social") as objeto, \
            patch.object(flask_app, "generar_estatutos") as comercial:
        resp = client.post("/api/generate", json=data)
        objeto.assert_not_called()
        comercial.assert_not_called()
    assert resp.status_code == 200, resp.get_data(as_text=True)[:600]
    assert "_Familia.zip" in resp.headers["Content-Disposition"]
    with zipfile.ZipFile(io.BytesIO(resp.data)) as zf:
        def contenido(fragmento):
            nombre = next(n for n in zf.namelist() if fragmento in n)
            return io.BytesIO(zf.read(nombre))

        txt = "\n".join(p.text for p in Document(contenido("Estatutos.docx")).paragraphs)
        assert "precautelación" in txt
        assert "No se pacta opción estatutaria de compra" in txt
        assert "Texto que debe ignorarse" not in txt
        assert "{{" not in txt and "[[" not in txt
        fam = " ".join(p.extract_text() or "" for p in
                       PdfReader(contenido("Formato_Empresa_Familiar")).pages)
        assert "Persona Test Uno" in fam and "Persona Test Dos" in fam
        assert "Padre" in fam and "Hija" in fam
        assert "6.000" in fam and "3.000" in fam
        assert "Holding Prueba" not in fam
        campos = PdfReader(contenido("Formulario_RUES")).get_fields()
        for indices, esperado in [((164, 165, 166, 167), "7010"),
                                  ((169, 170, 171, 172), "6201"),
                                  ((201, 202, 203, 204), "7010")]:
            assert "".join(str(campos[f"Cuadro de texto 2_{i}"].get("/V", ""))
                           for i in indices) == esperado
        assert str(campos["Casilla 1_37"].get("/V")) not in ("", "None", "/Off")
        assert str(campos["Casilla 1_38"].get("/V")) in ("", "None", "/Off")
        for fragmento in ("Otras_Entidades", "Responsabilidades_Tributarias",
                          "Emprendimiento_Social", "Grupo_Etnico", "Situacion_Control", "Soportes/"):
            assert any(fragmento in n for n in zf.namelist()), fragmento
    print("OK paquete familia: estatutos, anexos, núcleo familiar y CIIU del RUES")


def test_validaciones_familia(client):
    for parentesco in (None, "", "   "):
        data = payload_familia()
        data["accionistas"][0]["parentesco"] = parentesco
        resp = client.post("/api/generate", json=data)
        assert resp.status_code == 400
        assert "Persona Test Uno" in resp.get_json()["error"]
        assert "parentesco" in resp.get_json()["error"]
    # Se rechazan antes de exigir la tarjeta de abogado, incluso vacíos.
    for campo in ("disposiciones", "modelo_propio"):
        for valor in ({}, {"operaciones": [{"texto": "No permitido"}]}):
            resp = client.post("/api/generate", json=dict(payload_familia(), **{campo: valor}))
            assert resp.status_code == 400
            assert "no está" in resp.get_json()["error"]
    data = payload_familia()
    data["accionistas"][0]["porcentaje"] = 50
    resp = client.post("/api/generate", json=data)
    assert resp.status_code == 400 and "100 %" in resp.get_json()["error"]
    with patch.object(flask_app, "renderizar", side_effect=flask_app.PlantillaError("Falta un campo")):
        resp = client.post("/api/generate", json=payload_familia())
        assert resp.status_code == 400 and resp.get_json()["error"] == "Falta un campo"
    # La actividad secundaria sigue sometida a la validación comercial.
    with patch.object(flask_app, "validar_ciiu", return_value=(False, ["Secundario rechazado"], {})) as validar:
        resp = client.post("/api/generate", json=dict(payload_familia(), ciiu_code_sec="6201"))
        assert resp.status_code == 400
        assert validar.call_args.args[0] == ["7010", "6201"]
        assert validar.call_args.args[2] == flask_app.familia.OBJETO_PRINCIPAL
    print("OK validaciones familia: parentesco, opciones no disponibles y errores de contexto/plantilla")


def test_solo_estatutos():
    with flask_app.app.test_request_context():
        ruta = flask_app._generar_paquete(payload_familia(), {}, solo_estatutos=True)
        assert isinstance(ruta, str) and os.path.isfile(ruta)
        assert "precautelación" in "\n".join(p.text for p in Document(ruta).paragraphs)
    print("OK familia solo_estatutos")


if __name__ == "__main__":
    cli = _cliente()
    test_paquete_familia(cli)
    test_validaciones_familia(cli)
    test_solo_estatutos()
    print("OK test_generar_familia")
