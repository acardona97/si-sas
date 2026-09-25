# -*- coding: utf-8 -*-
"""Modelo del abogado: formato, ubicaciones, autorización y paquete final.

Ejecutar: python test_modelo_propio.py
"""
import io
import json
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from test_paquete import _cliente, PAYLOAD_B
import app as flask_app
from processors import modelo_propio as mp
from processors.estatutos import generar_estatutos
from processors.disposiciones import indexar

BASE = Path(__file__).resolve().parent


def crear_modelo(path):
    doc = Document()
    textos = [
        'Estatutos de EJEMPLO ANTERIOR S.A.S.',
        'Artículo 1. Domicilio: Cali.',
        'Artículo 2. Capital autorizado: $5.000.000.',
        'Artículo 3. Objeto: Venta de zapatos antiguos.',
        'Comparece ACCIONISTA ANTERIOR, C.C. 999.',
        'Artículo 4. Administración.',
        'Representante: GERENTE ANTERIOR, documento 888, expedido en Cali.',
        'Artículo Primero Transitorio. Nombramientos iniciales.',
        'Artículo 5. Facultades del representante legal.',
        'Artículo 6. Suscripción del capital.',
        'Los constituyentes suscriben estos estatutos.',
    ]
    for i, texto in enumerate(textos):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(.25)
        p.paragraph_format.space_after = Pt(7 + i)
        # El nombre y el objeto están partidos en runs con distinto énfasis.
        corte = max(1, len(texto) // 2)
        for j, parte in enumerate([texto[:corte], texto[corte:]]):
            r = p.add_run(parte)
            r.font.name = 'Arial'
            r.font.size = Pt(12)
            r.italic = bool(j)
    doc.save(path)
    return doc


def propuesta():
    fragmentos = [
        (0, 'EJEMPLO ANTERIOR S.A.S.', 'NOMBRE_SOCIEDAD'),
        (1, 'Cali', 'DOMICILIO_MUNICIPIO'),
        (2, '$5.000.000', 'CAPITAL_AUTORIZADO_MONTO_LETRAS_Y_CIFRAS'),
        (3, 'Venta de zapatos antiguos.', 'OBJETO_SOCIAL_DESARROLLADO'),
        (4, 'ACCIONISTA ANTERIOR, C.C. 999.', 'BLOQUE_COMPARECENCIA_ACCIONISTAS'),
        (6, 'GERENTE ANTERIOR', 'RL_PRINCIPAL_NOMBRE'),
        (6, '888', 'RL_PRINCIPAL_NUM_DOC'),
        (6, 'Cali', 'RL_PRINCIPAL_CIUDAD_EXPEDICION'),
    ]
    return {'reemplazos': [{'parrafo': i, 'buscar': texto, 'reemplazar_por_token': '{{' + token + '}}'}
                           for i, texto, token in fragmentos],
            'anclas': dict(zip(mp.ANCLAS, [7, 10, 9, 8])), 'no_ubicados': []}


def xml_propiedades(elem, nombre):
    p = elem.find(qn(nombre))
    return p.xml if p is not None else None


def assert_formato(original, resultado):
    # Cada párrafo original conserva sus propiedades y todos sus runs.
    originales = list(original.paragraphs)
    for i, p in enumerate(originales):
        referencia = xml_propiedades(p._p, 'w:pPr')
        candidatos = [q for q in resultado.paragraphs
                      if xml_propiedades(q._p, 'w:pPr') == referencia and len(q.runs) == len(p.runs)]
        assert candidatos, f'Párrafo {i} perdió su formato'
        q = candidatos[0]
        assert [xml_propiedades(r._r, 'w:rPr') for r in p.runs] == [
            xml_propiedades(r._r, 'w:rPr') for r in q.runs]
    for p in resultado.paragraphs:
        for r in p.runs:
            assert r.font.name == 'Arial' and r.font.size == Pt(12)


def datos():
    return dict(deepcopy(PAYLOAD_B), fecha='2026-09-25', capital_autorizado=9000000,
                capital_suscrito=1000000, capital_pagado=1000000,
                junta_directiva={'principales': [{'nombre': 'Pedro Nuevo', 'tipo_doc': 'CC', 'id_num': '123'}]},
                revisor_fiscal={'nombre': 'Luis Contador', 'tipo_doc': 'CC', 'id_num': '456', 'tarjeta_profesional': '321-T'},
                limitaciones_rl={'tiene_limitaciones': True, 'limita_cuantia': True,
                                 'cuantia_smmlv': '200', 'organo_cuantia': 'asamblea'})


def test_documento(carpeta):
    fuente, tokens, salida = [carpeta / n for n in ('modelo.docx', 'tokens.docx', 'estatutos.docx')]
    original = crear_modelo(fuente)
    assert mp.validar_docx(fuente.read_bytes())
    with patch.object(mp, '_llamar_sonnet', return_value=propuesta()) as sonnet:
        ubicacion = mp.ubicar_campos(fuente)
        assert sonnet.call_count == 1
    assert not ubicacion['usa_tokens_propios']
    assert all(aviso.startswith('Sin ubicación:') for aviso in ubicacion['no_ubicados'])
    mp.tokenizar(fuente, ubicacion, tokens)
    tokenizado = Document(tokens)
    assert '{{NOMBRE_SOCIEDAD}}' in tokenizado.paragraphs[0].text
    assert_formato(original, tokenizado)
    d = dict(datos(), anclas_modelo_propio=ubicacion['anclas'])
    generar_estatutos(d, tokens, salida)
    generado = Document(salida)
    texto = '\n'.join(p.text for p in generado.paragraphs)
    assert '{{' not in texto
    for viejo in ['EJEMPLO ANTERIOR', 'Cali', '$5.000.000', 'zapatos antiguos', 'ACCIONISTA ANTERIOR', 'GERENTE ANTERIOR', '888', '999']:
        assert viejo not in texto, viejo
    for nuevo in ['SOLO UNO S.A.S.', 'Medellín', '9.000.000', 'Consultoría empresarial',
                  'ANA RESTREPO GÓMEZ', 'PEDRO NUEVO', 'LUIS CONTADOR', '200 salarios']:
        assert nuevo in texto, nuevo
    indice = next(i for i, p in enumerate(generado.paragraphs) if 'Primero Transitorio' in p.text)
    assert 'PEDRO NUEVO' in generado.paragraphs[indice + 1].text
    assert generado.tables and 'ANA RESTREPO GÓMEZ' in generado.tables[0].cell(1, 0).text
    assert_formato(original, generado)

    # Sin anclas no se adivinan inserciones ni se altera la estructura base.
    generar_estatutos(dict(datos(), anclas_modelo_propio={}), tokens, salida)
    assert len(Document(salida).paragraphs) == len(original.paragraphs)

    # Una plantilla ya tokenizada no necesita llamada a Claude.
    with patch.object(mp, '_llamar_sonnet', side_effect=AssertionError('No debe llamar IA')):
        assert mp.ubicar_campos(tokens)['usa_tokens_propios']

    # Disposiciones aprobadas siguen aplicándose encima del modelo propio.
    generar_estatutos(d, tokens, salida)
    entrada = indexar(Document(salida))[8]
    op = {'tipo': 'insertar', 'forma': 'inciso', 'parrafo': entrada['i'],
          'ancla': entrada['texto'], 'texto': 'Regla adicional aprobada.',
          'articulo': entrada['articulo'], 'disposicion': 1, 'motivo': 'Prueba'}
    for cambios in [False, True]:
        generar_estatutos(dict(d, disposiciones=[op], disposiciones_con_cambios=cambios), tokens, salida)
        final = Document(salida)
        assert 'Regla adicional aprobada.' in ''.join(t.text or '' for t in final.element.iter(qn('w:t')))
        assert bool(list(final.element.iter(qn('w:ins')))) == cambios
    print('OK modelo: tokens, datos, formato, anclas y disposiciones')


def test_validacion(carpeta):
    path = carpeta / 'modelo.docx'
    for contenido in [b'%PDF-invalido', b'no es word', b'x' * (mp.MAX_BYTES + 1)]:
        try:
            mp.validar_docx(contenido)
            raise AssertionError('Debió rechazar el archivo')
        except mp.ModeloPropioError:
            pass
    malas = propuesta()
    malas['reemplazos'].extend([
        {'parrafo': -1, 'buscar': 'Cali', 'reemplazar_por_token': '{{NOMBRE_SOCIEDAD}}'},
        {'parrafo': 1, 'buscar': 'Bogotá', 'reemplazar_por_token': '{{NOMBRE_SOCIEDAD}}'},
        {'parrafo': 1, 'buscar': 'Cali', 'reemplazar_por_token': '{{INVENTADO}}'},
        {'parrafo': 0, 'buscar': 'ANTERIOR', 'reemplazar_por_token': '{{NOMBRE_SOCIEDAD}}'},
    ])
    malas['anclas']['ancla_firmas'] = 9999
    with patch.object(mp, '_llamar_sonnet', return_value=malas):
        resultado = mp.ubicar_campos(path)
    assert len([aviso for aviso in resultado['no_ubicados'] if not aviso.startswith('Sin ubicación:')]) == 5
    assert resultado['anclas']['ancla_firmas'] is None
    try:
        mp.tokenizar(path, malas, carpeta / 'invalido.docx')
        raise AssertionError('Debió rechazar operaciones adulteradas')
    except mp.ModeloPropioError:
        pass
    print('OK validación de archivos y operaciones')


def test_endpoint(carpeta):
    client = _cliente()
    contenido = (carpeta / 'modelo.docx').read_bytes()

    def preview(nombre='modelo.docx', archivo=contenido, payload=None):
        return client.post('/api/modelo-propio/preview', data={
            'modelo': (io.BytesIO(archivo), nombre),
            'payload': json.dumps(PAYLOAD_B if payload is None else payload)},
            content_type='multipart/form-data')

    assert preview().status_code == 403
    with client.session_transaction() as ses:
        ses['tp_abogado'] = {'numero_tarjeta': '123456'}
    for nombre, contenido_malo in [('modelo.pdf', contenido), ('modelo.doc', contenido), ('modelo.docx', b'error')]:
        assert preview(nombre, contenido_malo).status_code == 400
    assert preview(payload=[]).status_code == 400
    with patch.object(mp, '_llamar_sonnet', return_value=propuesta()):
        respuesta = preview()
    assert respuesta.status_code == 200, respuesta.get_data(as_text=True)
    aprobado = respuesta.get_json()
    assert set(aprobado) == {'modelo_id', 'reemplazos', 'anclas', 'no_ubicados', 'usa_tokens_propios'}
    for identificador in ['../modelo', 'f' * 32, None]:
        r = client.post('/api/generate', json=dict(PAYLOAD_B, modelo_propio={'modelo_id': identificador}))
        assert r.status_code == 400, r.get_data(as_text=True)
    alterado = deepcopy(aprobado)
    alterado['reemplazos'][0]['fragmento'] = 'No existe este nombre'
    with patch.object(flask_app, 'generar_objeto_social', return_value='Consultoría empresarial.'):
        assert client.post('/api/generate', json=dict(PAYLOAD_B, modelo_propio=alterado)).status_code == 400
        payload = dict(PAYLOAD_B, modelo_propio=aprobado)
        respuesta = client.post('/api/generate', json=payload)
    assert respuesta.status_code == 200, respuesta.get_data(as_text=True)[:500]
    with zipfile.ZipFile(io.BytesIO(respuesta.data)) as z:
        nombre = next(n for n in z.namelist() if n.endswith('_Estatutos.docx'))
        doc = Document(io.BytesIO(z.read(nombre)))
        assert 'SOLO UNO S.A.S.' in doc.paragraphs[0].text
        assert 'EJEMPLO ANTERIOR' not in '\n'.join(p.text for p in doc.paragraphs)
    with client.session_transaction() as ses:
        assert 'tp_abogado' not in ses
        ses['tp_abogado'] = {'numero_tarjeta': '123456'}
        ses['user_id'] = -123
    with flask_app.app.test_request_context('/'):
        flask_app.session['user_id'] = -123
        try:
            flask_app._ruta_modelo_propio(aprobado)
            raise AssertionError('No debe permitir modelos ajenos')
        except mp.ModeloPropioError:
            pass
    print('OK preview: 403/200/400; generación y propietario')


if __name__ == '__main__':
    with tempfile.TemporaryDirectory(prefix='_test_modelo_', dir=BASE / 'output') as carpeta:
        carpeta = Path(carpeta)
        test_documento(carpeta)
        test_validacion(carpeta)
        test_endpoint(carpeta)
    print('OK test_modelo_propio')
