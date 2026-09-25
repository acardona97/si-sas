"""Disposiciones especiales: aplicación determinista sobre el Word (sin IA)."""
import os
import re
import sys
import tempfile
from copy import deepcopy
from datetime import date

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "src"))
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from processors.estatutos import generar_estatutos  # noqa: E402
from processors.disposiciones import indexar, DisposicionError  # noqa: E402

TMPL = os.path.join(BASE, "plantillas", "estatutos_template.docx")
p = lambda n: {"nombre": f"Persona Test {n}", "tipo_doc": "CC", "id_num": f"9999999{n}"}  # noqa: E731
DATA = {
    "nombre_sas": "PRUEBA S.A.S.", "fecha": date(2026, 9, 25),
    "municipio": "Medellín", "departamento": "Antioquia",
    "accionistas": [{"tipo": "natural", "nombre": "Persona Test", "id_num": "99999999",
                     "expedicion": "Medellín", "domicilio": "Medellín", "genero": "M", "porcentaje": 100}],
    "objeto_social": "cualquier actividad comercial o civil lícita",
    "capital_suscrito": 1_000_000, "capital_pagado": 1_000_000,
    "rl_principales": [{"nombre": "Persona Test", "cc": "99999999", "tipo_doc": "CC", "expedicion": "Medellín"}],
    "rl_suplentes": [],
    "junta_directiva": {"principales": [p(1), p(2), p(3)], "suplentes": []},
    "revisor_fiscal": {"tipo": "natural", "nombre": "Persona Test 8", "id_num": "3", "tarjeta_profesional": "1-T"},
}


def generar(extra=None):
    tmp = tempfile.mkdtemp()
    out = os.path.join(tmp, "e.docx")
    generar_estatutos(dict(DATA, **(extra or {})), TMPL, out)
    return Document(out)


base = generar()
ent = indexar(base)
buscar = lambda inicio: next(e for e in ent if e["texto"].startswith(inicio))  # noqa: E731
reuniones = buscar("Reuniones, quórum y mayorías de la junta directiva.")
revisoria = buscar("Nombramiento y período de la revisoría fiscal.")
funciones_rf = buscar("Funciones. La revisoría fiscal")
rev_par = buscar("Parágrafo. El revisor fiscal podrá intervenir")
quorum = buscar("Quórum. Salvo disposición")
assert reuniones["tipo"] == "articulo" and funciones_rf["articulo"] == rev_par["articulo"]

op = lambda tipo, forma, e, texto="": {  # noqa: E731
    "tipo": tipo, "forma": forma, "parrafo": e["i"], "ancla": e["texto"],
    "texto": texto, "disposicion": 1, "motivo": "prueba", "articulo": e["articulo"]}
OPS = [
    op("insertar", "paragrafo", reuniones,
       "Las reuniones de la junta directiva podrán celebrarse en cualquier lugar del país."),
    op("insertar", "paragrafo", funciones_rf,
       "El revisor fiscal presentará informes trimestrales a la junta directiva."),
    op("insertar", "inciso", revisoria,
       "La asamblea general de accionistas fijará los honorarios del revisor fiscal."),
    op("reemplazar", "ninguna", reuniones,
       "La junta directiva se reunirá ordinariamente por lo menos una (1) vez al mes."),
    op("reemplazar", "ninguna", quorum, "La asamblea podrá deliberar con cualquier número de accionistas."),
]

limpio = generar({"disposiciones": OPS})
txt = "\n".join(x.text for x in limpio.paragraphs)
# Parágrafo nuevo en artículo sin parágrafos → "Parágrafo."
assert "Parágrafo. Las reuniones de la junta directiva podrán celebrarse" in txt
# Artículo que ya tenía "Parágrafo." → pasan a primero y segundo
assert "Parágrafo primero. El revisor fiscal podrá intervenir" in txt, txt[-3000:]
assert "Parágrafo segundo. El revisor fiscal presentará informes trimestrales" in txt
# El título del artículo se conserva al reemplazar el cuerpo
assert "Reuniones, quórum y mayorías de la junta directiva. La junta directiva se reunirá ordinariamente por lo menos una (1) vez al mes." in txt
assert "cada tres (3) meses" not in txt
# La puntuación tras el título en negrilla vive en el run del cuerpo y se conserva
assert "Quórum. La asamblea podrá deliberar con cualquier número de accionistas." in txt
# El inciso va inmediatamente después del primer párrafo del artículo
ps = [x.text for x in limpio.paragraphs if x.text.strip()]
k = next(i for i, t in enumerate(ps) if t.startswith("Nombramiento y período de la revisoría fiscal."))
assert ps[k + 1].startswith("La asamblea general de accionistas fijará los honorarios"), ps[k + 1]
# No agrega líneas en blanco dobles (las del bloque de firmas y alrededor de
# la tabla de junta ya existen en el documento base)
def dobles(doc):
    v = [not x.text.strip() for x in doc.paragraphs]
    return sum(a and b for a, b in zip(v, v[1:]))
assert dobles(limpio) == dobles(base), "quedaron dos líneas en blanco seguidas"

# Formato: todo párrafo/run nuevo usa una firma que ya existe en el documento base
def firmas(doc):
    pp, rr = set(), set()
    for x in doc.paragraphs:
        ppr = x._p.find(qn("w:pPr"))
        pp.add(re.sub(r"<w:rPr>.*?</w:rPr>|<w:rPr/>", "", str(ppr.xml) if ppr is not None else "", flags=re.S))
        for r in x._p.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            rr.add(str(rpr.xml) if rpr is not None else "")
    return pp, rr

bp, br = firmas(base)
lp, lr = firmas(limpio)
assert lp <= bp, "párrafos con formato ajeno a la plantilla"
assert lr <= br, "runs con formato ajeno a la plantilla"

# Control de cambios: aceptar todo reproduce la versión limpia
cambios = generar({"disposiciones": OPS, "disposiciones_con_cambios": True})
body = cambios.element.body
assert body.findall(".//" + qn("w:ins")) and body.findall(".//" + qn("w:del"))
aceptado = deepcopy(body)
for d in aceptado.findall(".//" + qn("w:del")):
    padre = d.getparent()
    if padre.tag == qn("w:rPr"):   # marca de párrafo eliminado
        par = padre.getparent().getparent()
        par.getparent().remove(par)
    else:
        padre.remove(d)
for i in aceptado.findall(".//" + qn("w:ins")):
    if i.getparent().tag == qn("w:rPr"):
        i.getparent().remove(i)
    else:
        for r in list(i):
            i.addprevious(r)
        i.getparent().remove(i)
txt_aceptado = "\n".join("".join(t.text or "" for t in par.iter(qn("w:t")))
                         for par in aceptado.findall(qn("w:p")))
assert [t for t in txt_aceptado.split("\n") if t.strip()] == [t for t in txt.split("\n") if t.strip()]

# Si los estatutos cambiaron desde la vista previa, se detiene
malo = dict(OPS[0], ancla="texto que ya no existe", parrafo=3)
try:
    generar({"disposiciones": [malo]})
    raise AssertionError("debía detenerse")
except DisposicionError:
    pass

print("OK test_disposiciones")
